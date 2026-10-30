"""Spreadsheet conversion and inspection."""

from __future__ import annotations

import csv
from pathlib import Path

from .base_converter import ConversionResult, Converter
from .utils.exceptions import ConversionError


class SpreadsheetConverter(Converter):
    def xlsx_to_csv(self, source: str | Path, destination: str | Path, sheet: str | None = None) -> ConversionResult:
        path = self.validate_input(source)
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_name = sheet or workbook.sheetnames[0]
            if sheet_name not in workbook.sheetnames:
                raise ConversionError(f"Worksheet not found: {sheet_name}")
            output = Path(destination)
            with output.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                for row in workbook[sheet_name].iter_rows(values_only=True):
                    writer.writerow(["" if value is None else value for value in row])
            workbook.close()
            return self.result(path, output)
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError(f"Unable to read XLSX: {exc}") from exc

    def csv_to_xlsx(self, source: str | Path, destination: str | Path) -> ConversionResult:
        path = self.validate_input(source)
        try:
            from openpyxl import Workbook
            workbook = Workbook()
            sheet = workbook.active
            with path.open(newline="", encoding="utf-8-sig") as handle:
                for row in csv.reader(handle):
                    sheet.append(row)
            output = Path(destination)
            workbook.save(output)
            return self.result(path, output)
        except Exception as exc:
            raise ConversionError(f"Unable to create XLSX: {exc}") from exc

    def inspect(self, source: str | Path, sheet: str | None = None) -> dict[str, object]:
        path = self.validate_input(source)
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=True)
            selected = sheet or workbook.sheetnames[0]
            rows = list(workbook[selected].iter_rows(values_only=True))
            workbook.close()
            return {"sheets": workbook.sheetnames, "selected_sheet": selected, "rows": len(rows), "columns": max((len(row) for row in rows), default=0)}
        except Exception as exc:
            raise ConversionError(f"Unable to inspect XLSX: {exc}") from exc

    @staticmethod
    def validate_csv(source: str | Path) -> tuple[bool, str]:
        path = Path(source)
        try:
            with path.open(newline="", encoding="utf-8-sig") as handle:
                rows = list(csv.reader(handle))
            if not rows:
                return False, "CSV is empty."
            width = len(rows[0])
            if any(len(row) != width for row in rows[1:]):
                return False, "CSV rows have inconsistent column counts."
            return True, "CSV is valid."
        except Exception as exc:
            return False, str(exc)
