"""Document conversion helpers using local libraries and optional Pandoc."""

from __future__ import annotations

from pathlib import Path

from .base_converter import ConversionResult, Converter
from .utils.exceptions import ConversionError, UnsupportedConversionError


class DocumentConverter(Converter):
    def docx_to_text(self, source: str | Path) -> str:
        path = self.validate_input(source)
        try:
            from docx import Document
            document = Document(path)
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        except Exception as exc:
            raise ConversionError(f"Unable to read DOCX: {exc}") from exc

    def text_to_docx(self, source: str | Path, destination: str | Path) -> ConversionResult:
        path = self.validate_input(source)
        try:
            from docx import Document
            document = Document()
            for line in Path(path).read_text(encoding="utf-8", errors="replace").splitlines():
                document.add_paragraph(line)
            output = Path(destination)
            document.save(output)
            return self.result(path, output)
        except Exception as exc:
            raise ConversionError(f"Unable to create DOCX: {exc}") from exc

    def markdown_to_html(self, source: str | Path, destination: str | Path) -> ConversionResult:
        path = self.validate_input(source)
        text = path.read_text(encoding="utf-8", errors="replace")
        try:
            import pypandoc
            html = pypandoc.convert_text(text, "html", format="md")
        except Exception:
            html = f"<html><body><pre>{_escape(text)}</pre></body></html>"
        output = Path(destination)
        output.write_text(html, encoding="utf-8")
        return self.result(path, output)

    def text_to_pdf(self, source: str | Path, destination: str | Path) -> ConversionResult:
        path = self.validate_input(source)
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen.canvas import Canvas
            output = Path(destination)
            canvas = Canvas(str(output), pagesize=letter)
            y = 750
            for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
                canvas.drawString(40, y, line[:120])
                y -= 14
                if y < 40:
                    canvas.showPage()
                    y = 750
            canvas.save()
            return self.result(path, output)
        except Exception as exc:
            raise ConversionError(f"Unable to create PDF: {exc}") from exc


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
